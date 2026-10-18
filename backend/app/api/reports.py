import csv
import io
from collections import Counter
from datetime import datetime, timedelta, timezone
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import current_user
from app.models import AuditEvent, Invoice, User
from app.core.config import settings
from app.core.dependencies import require_roles
from app.models import Supplier
from app.services.scheduler import scan_inbox_folder
from pathlib import Path

router = APIRouter(prefix="/reports", tags=["Module 6 · Reporting & Monitor"])


@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), _: User = Depends(current_user)):
    total = db.query(func.count(Invoice.id)).scalar() or 0
    counts = dict(db.query(Invoice.status, func.count(Invoice.id)).group_by(Invoice.status).all())
    recent = db.query(Invoice).order_by(Invoice.created_at.desc()).limit(6).all()
    today = datetime.now(timezone.utc).date()
    days = [today - timedelta(days=n) for n in range(6, -1, -1)]
    weekly = {day.isoformat(): {"matched": 0, "mismatch": 0, "review": 0} for day in days}
    warnings = Counter()
    for row in db.query(Invoice).all():
        created_day = row.created_at.date() if row.created_at else None
        bucket = weekly.get(created_day.isoformat()) if created_day else None
        if bucket is not None:
            key = "matched" if row.status in {"matched", "approved"} else "mismatch" if row.status == "mismatch" else "review"
            bucket[key] += 1
        for warning in row.warnings or []:
            warnings[warning.get("code", "OTHER")] += 1
    supplier_names = {s.tax_code: s.name for s in db.query(Supplier).all()}
    supplier_counts = Counter(x.seller_tax_code for x in db.query(Invoice).filter(Invoice.seller_tax_code != "").all())
    top_suppliers = [{"tax_code": code, "name": supplier_names.get(code, "Chưa có trong sổ NCC"), "invoices": n}
                     for code, n in supplier_counts.most_common(5)]
    return {"total": total, "by_status": counts,
            "needs_review": counts.get("needs_review", 0) + counts.get("mismatch", 0),
            "approved": counts.get("approved", 0), "weekly": [{"day": d, **weekly[d]} for d in weekly],
            "error_causes": [{"code": code, "count": n} for code, n in warnings.most_common(6)],
            "top_suppliers": top_suppliers, "supplier_count": db.query(func.count(Supplier.id)).scalar() or 0,
            "recent": [{"id": x.id, "name": x.original_name, "status": x.status,
            "amount": x.total_amount, "created_at": x.created_at.isoformat()} for x in recent]}


@router.get("/worker")
def worker_status(_: User = Depends(current_user)):
    inbox = Path(settings.inbox_scan_dir) if settings.inbox_scan_dir else None
    active = bool(inbox and inbox.is_dir())
    return {"active": active, "mode": "local_folder_poll" if active else "not_configured",
            "message": "Đang theo dõi thư mục inbox mỗi phút" if active else "Chưa cấu hình thư mục inbox"}


@router.post("/scan-now")
def scan_now(_: User = Depends(require_roles("admin"))):
    inbox = Path(settings.inbox_scan_dir) if settings.inbox_scan_dir else None
    if not inbox or not inbox.is_dir():
        return {"active": False, "added": 0, "message": "Chưa cấu hình thư mục inbox; không có file nào được quét."}
    count = scan_inbox_folder()
    return {"active": True, "added": count, "message": f"Đã thêm {count} hóa đơn mới từ thư mục inbox."}


@router.get("/audit")
def audit_log(limit: int = 100, db: Session = Depends(get_db), _: User = Depends(current_user)):
    rows = db.query(AuditEvent).order_by(AuditEvent.created_at.desc()).limit(min(limit, 500)).all()
    return [{"id": x.id, "actor": x.actor_email, "action": x.action, "entity": f"{x.entity_type}:{x.entity_id}",
             "detail": x.detail, "created_at": x.created_at.isoformat()} for x in rows]


@router.get("/export.csv")
def export_csv(db: Session = Depends(get_db), _: User = Depends(current_user)):
    buffer = io.StringIO(); writer = csv.writer(buffer)
    writer.writerow(["id", "seller_tax_code", "buyer_tax_code", "invoice_number", "invoice_date", "total_amount", "status", "po_number"])
    for x in db.query(Invoice).order_by(Invoice.created_at.desc()):
        writer.writerow([x.id, x.seller_tax_code, x.buyer_tax_code, x.invoice_number, x.invoice_date, x.total_amount, x.status, x.po_number])
    return StreamingResponse(iter([buffer.getvalue()]), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": "attachment; filename=invoice-report.csv"})


@router.get("/export.xlsx")
def export_xlsx(db: Session = Depends(get_db), _: User = Depends(current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "Invoices"
    headers = ["ID", "MST bên bán", "MST bên mua", "Số hóa đơn", "Ngày", "Tổng tiền", "Trạng thái", "PO"]
    ws.append(headers)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="243B53")
    for x in db.query(Invoice).order_by(Invoice.created_at.desc()):
        ws.append([x.id, x.seller_tax_code, x.buyer_tax_code, x.invoice_number, x.invoice_date, x.total_amount, x.status, x.po_number])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = min(34, max(12, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[get_column_letter(col[0].column)].width = width
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=invoice-report.xlsx"})
