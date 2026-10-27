import csv
import io
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import current_user, require_roles
from app.models import AuditEvent, Customer, Employee, GoodsReceiptItem, InvoiceItem, Product, ProductCategory, PurchaseOrderItem, User

router = APIRouter(prefix="/master-data", tags=["Master data"])
EDIT_ROLES = ("admin", "chief_accountant")


class CategoryIn(BaseModel):
    code: str = Field(min_length=1, max_length=80)
    name: str = Field(min_length=1, max_length=255)
    description: str = ""


class CategoryPatch(BaseModel):
    name: str | None = Field(default=None, min_length=1, max_length=255)
    description: str | None = None


class ProductIn(BaseModel):
    code: str = Field(min_length=1, max_length=80)
    name: str = Field(min_length=1, max_length=500)
    category_code: str | None = None
    unit: str = ""
    cost_price: float = Field(default=0, ge=0)
    selling_price: float = Field(default=0, ge=0)
    stock_qty: float = Field(default=0, ge=0)


class ProductPatch(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str | None = Field(default=None, min_length=1, max_length=500)
    category_code: str | None = None
    unit: str | None = None
    cost_price: float | None = Field(default=None, ge=0)
    selling_price: float | None = Field(default=None, ge=0)
    stock_qty: float | None = Field(default=None, ge=0)


class CustomerIn(BaseModel):
    code: str = Field(min_length=1, max_length=80)
    name: str = Field(min_length=1, max_length=255)
    phone: str = ""
    address: str = ""
    email: str = ""


class CustomerPatch(BaseModel):
    name: str | None = Field(default=None, min_length=1, max_length=255)
    phone: str | None = None
    address: str | None = None
    email: str | None = None


class EmployeeIn(BaseModel):
    code: str = Field(min_length=1, max_length=80)
    full_name: str = Field(min_length=1, max_length=255)
    role: str = "Kế toán viên"
    phone: str = ""
    email: str = ""


class EmployeePatch(BaseModel):
    full_name: str | None = Field(default=None, min_length=1, max_length=255)
    role: str | None = None
    phone: str | None = None
    email: str | None = None


def _audit(db: Session, user: User, action: str, kind: str, code: str, detail: str = "") -> None:
    db.add(AuditEvent(actor_user_id=getattr(user, "id", None), actor_email=user.email,
                      action=action, entity_type=kind, entity_id=code, detail=detail))


def _commit(db: Session, duplicate_message: str):
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, duplicate_message) from exc


def _category(row: ProductCategory) -> dict:
    return {"id": row.id, "code": row.code, "name": row.name, "description": row.description or ""}


def _product(row: Product) -> dict:
    return {"id": row.id, "code": row.sku or f"SP-{row.id}", "name": row.product_name,
            "category_code": row.category.code if row.category else "", "unit": row.unit or "",
            "cost_price": float(row.cost_price or 0), "selling_price": float(row.selling_price or 0),
            "stock_qty": float(row.stock_qty or 0)}


def _customer(row: Customer) -> dict:
    return {"id": row.id, "code": row.code, "name": row.name, "phone": row.phone or "",
            "address": row.address or "", "email": row.email or ""}


def _employee(row: Employee) -> dict:
    return {"id": row.id, "code": row.code, "full_name": row.full_name, "role": row.role,
            "phone": row.phone or "", "email": row.email or ""}


async def _read_import_rows(file: UploadFile) -> list[dict]:
    raw = await file.read()
    name = (file.filename or "").lower()
    if name.endswith(".csv"):
        return list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values, ())]
        return [dict(zip(headers, row)) for row in values]
    raise HTTPException(415, "Chỉ nhận tệp CSV hoặc XLSX")


@router.post("/import/{kind}")
async def import_master_data(
    kind: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*EDIT_ROLES)),
):
    models = {
        "product-categories": (ProductCategory, "code"),
        "products": (Product, "sku"),
        "customers": (Customer, "code"),
        "employees": (Employee, "code"),
    }
    if kind not in models:
        raise HTTPException(404, "Loại danh mục không hỗ trợ import")
    rows = await _read_import_rows(file)
    created = skipped = 0
    for values in rows:
        values = {str(key or "").strip(): value for key, value in values.items()}
        model, key_name = models[kind]
        try:
            if kind == "product-categories":
                key, name = str(values.get("code", "")).strip(), str(values.get("name", "")).strip()
                if not key or not name or db.query(ProductCategory).filter_by(code=key).first():
                    skipped += 1
                    continue
                row = ProductCategory(code=key, name=name, description=str(values.get("description", "") or ""))
            elif kind == "products":
                key, name = str(values.get("code", "")).strip(), str(values.get("name", "")).strip()
                category_code = str(values.get("category_code", "") or "").strip()
                category = db.query(ProductCategory).filter_by(code=category_code).first() if category_code else None
                if not key or not name or db.query(Product).filter_by(sku=key).first() or (category_code and category is None):
                    skipped += 1
                    continue
                row = Product(sku=key, product_name=name, category=category,
                              unit=str(values.get("unit", "") or ""),
                              cost_price=float(values.get("cost_price", 0) or 0),
                              selling_price=float(values.get("selling_price", 0) or 0),
                              stock_qty=float(values.get("stock_qty", 0) or 0))
            elif kind == "customers":
                key, name = str(values.get("code", "")).strip(), str(values.get("name", "")).strip()
                if not key or not name or db.query(Customer).filter_by(code=key).first():
                    skipped += 1
                    continue
                row = Customer(code=key, name=name, phone=str(values.get("phone", "") or ""),
                               address=str(values.get("address", "") or ""), email=str(values.get("email", "") or ""))
            else:
                key, name = str(values.get("code", "")).strip(), str(values.get("full_name", "")).strip()
                if not key or not name or db.query(Employee).filter_by(code=key).first():
                    skipped += 1
                    continue
                row = Employee(code=key, full_name=name, role=str(values.get("role", "Kế toán viên") or "Kế toán viên"),
                               phone=str(values.get("phone", "") or ""), email=str(values.get("email", "") or ""))
            db.add(row)
            created += 1
        except (TypeError, ValueError):
            skipped += 1
    db.flush()
    _audit(db, user, "master_data.import", kind, "bulk", f"created={created};skipped={skipped}")
    db.commit()
    return {"created": created, "skipped": skipped, "total_rows": len(rows)}


@router.get("/product-categories")
def list_categories(db: Session = Depends(get_db), _: User = Depends(current_user)):
    return [_category(row) for row in db.query(ProductCategory).order_by(ProductCategory.name)]


@router.post("/product-categories", status_code=201)
def create_category(body: CategoryIn, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = ProductCategory(**body.model_dump())
    db.add(row)
    _audit(db, user, "product_category.create", "product_category", row.code, row.name)
    _commit(db, "Mã loại sản phẩm đã tồn tại")
    db.refresh(row)
    return _category(row)


@router.patch("/product-categories/{code}")
def update_category(code: str, body: CategoryPatch, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = db.query(ProductCategory).filter_by(code=code).first()
    if row is None:
        raise HTTPException(404, "Không tìm thấy loại sản phẩm")
    values = body.model_dump(exclude_unset=True)
    if not values or any(value is None for value in values.values()):
        raise HTTPException(422, "Cần cung cấp ít nhất một giá trị hợp lệ")
    for key, value in values.items():
        setattr(row, key, value)
    _audit(db, user, "product_category.update", "product_category", row.code, ",".join(sorted(values)))
    _commit(db, "Không thể cập nhật loại sản phẩm")
    db.refresh(row)
    return _category(row)


@router.delete("/product-categories/{code}")
def delete_category(code: str, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = db.query(ProductCategory).filter_by(code=code).first()
    if row is None:
        raise HTTPException(404, "Không tìm thấy loại sản phẩm")
    if row.products:
        raise HTTPException(409, "Không thể xóa loại sản phẩm đang có sản phẩm tham chiếu")
    _audit(db, user, "product_category.delete", "product_category", row.code, row.name)
    db.delete(row)
    db.commit()
    return {"message": "Đã xóa loại sản phẩm"}


@router.get("/products")
def list_products(db: Session = Depends(get_db), _: User = Depends(current_user)):
    return [_product(row) for row in db.query(Product).order_by(Product.product_name)]


@router.post("/products", status_code=201)
def create_product(body: ProductIn, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    category = None
    if body.category_code:
        category = db.query(ProductCategory).filter_by(code=body.category_code).first()
        if category is None:
            raise HTTPException(422, "Không tìm thấy loại sản phẩm đã chọn")
    row = Product(sku=body.code, product_name=body.name, unit=body.unit, category=category,
                  cost_price=body.cost_price, selling_price=body.selling_price, stock_qty=body.stock_qty)
    db.add(row)
    _audit(db, user, "product.create", "product", body.code, body.name)
    _commit(db, "Mã sản phẩm đã tồn tại")
    db.refresh(row)
    return _product(row)


@router.patch("/products/{code}")
def update_product(code: str, body: ProductPatch, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = db.query(Product).filter_by(sku=code).first()
    if row is None:
        raise HTTPException(404, "Không tìm thấy sản phẩm")
    values = body.model_dump(exclude_unset=True)
    if not values or any(value is None for value in values.values()):
        raise HTTPException(422, "Cần cung cấp ít nhất một giá trị hợp lệ")
    if "category_code" in values:
        category_code = values.pop("category_code")
        row.category = db.query(ProductCategory).filter_by(code=category_code).first() if category_code else None
        if category_code and row.category is None:
            raise HTTPException(422, "Không tìm thấy loại sản phẩm đã chọn")
    if "name" in values:
        values["product_name"] = values.pop("name")
    for key, value in values.items():
        setattr(row, key, value)
    _audit(db, user, "product.update", "product", code, ",".join(sorted(body.model_fields_set)))
    _commit(db, "Không thể cập nhật sản phẩm")
    db.refresh(row)
    return _product(row)


@router.delete("/products/{code}")
def delete_product(code: str, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = db.query(Product).filter_by(sku=code).first()
    if row is None:
        raise HTTPException(404, "Không tìm thấy sản phẩm")
    used = any(db.query(model).filter_by(product_id=row.id).first() for model in (InvoiceItem, PurchaseOrderItem, GoodsReceiptItem))
    if used:
        raise HTTPException(409, "Không thể xóa sản phẩm đã được dùng trong chứng từ")
    _audit(db, user, "product.delete", "product", code, row.product_name)
    db.delete(row)
    db.commit()
    return {"message": "Đã xóa sản phẩm"}


@router.get("/customers")
def list_customers(db: Session = Depends(get_db), _: User = Depends(current_user)):
    return [_customer(row) for row in db.query(Customer).order_by(Customer.name)]


@router.post("/customers", status_code=201)
def create_customer(body: CustomerIn, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = Customer(**body.model_dump())
    db.add(row)
    _audit(db, user, "customer.create", "customer", row.code, row.name)
    _commit(db, "Mã khách hàng đã tồn tại")
    db.refresh(row)
    return _customer(row)


@router.patch("/customers/{code}")
def update_customer(code: str, body: CustomerPatch, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = db.query(Customer).filter_by(code=code).first()
    if row is None:
        raise HTTPException(404, "Không tìm thấy khách hàng")
    values = body.model_dump(exclude_unset=True)
    if not values or any(value is None for value in values.values()):
        raise HTTPException(422, "Cần cung cấp ít nhất một giá trị hợp lệ")
    for key, value in values.items():
        setattr(row, key, value)
    _audit(db, user, "customer.update", "customer", row.code, ",".join(sorted(values)))
    _commit(db, "Không thể cập nhật khách hàng")
    db.refresh(row)
    return _customer(row)


@router.delete("/customers/{code}")
def delete_customer(code: str, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = db.query(Customer).filter_by(code=code).first()
    if row is None:
        raise HTTPException(404, "Không tìm thấy khách hàng")
    _audit(db, user, "customer.delete", "customer", row.code, row.name)
    db.delete(row)
    db.commit()
    return {"message": "Đã xóa khách hàng"}


@router.get("/employees")
def list_employees(db: Session = Depends(get_db), _: User = Depends(current_user)):
    return [_employee(row) for row in db.query(Employee).order_by(Employee.full_name)]


@router.post("/employees", status_code=201)
def create_employee(body: EmployeeIn, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = Employee(**body.model_dump())
    db.add(row)
    _audit(db, user, "employee.create", "employee", row.code, row.full_name)
    _commit(db, "Mã nhân viên đã tồn tại")
    db.refresh(row)
    return _employee(row)


@router.patch("/employees/{code}")
def update_employee(code: str, body: EmployeePatch, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = db.query(Employee).filter_by(code=code).first()
    if row is None:
        raise HTTPException(404, "Không tìm thấy nhân viên")
    values = body.model_dump(exclude_unset=True)
    if not values or any(value is None for value in values.values()):
        raise HTTPException(422, "Cần cung cấp ít nhất một giá trị hợp lệ")
    for key, value in values.items():
        setattr(row, key, value)
    _audit(db, user, "employee.update", "employee", row.code, ",".join(sorted(values)))
    _commit(db, "Không thể cập nhật nhân viên")
    db.refresh(row)
    return _employee(row)


@router.delete("/employees/{code}")
def delete_employee(code: str, db: Session = Depends(get_db), user: User = Depends(require_roles(*EDIT_ROLES))):
    row = db.query(Employee).filter_by(code=code).first()
    if row is None:
        raise HTTPException(404, "Không tìm thấy nhân viên")
    _audit(db, user, "employee.delete", "employee", row.code, row.full_name)
    db.delete(row)
    db.commit()
    return {"message": "Đã xóa nhân viên"}
