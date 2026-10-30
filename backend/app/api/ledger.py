import csv
import io
import json
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import current_user, require_roles
from app.models import AuditEvent, GoodsReceipt, Invoice, PurchaseOrder, Supplier, User
from app.services.procurement_lines import sync_goods_receipt_lines, sync_purchase_order_lines
from app.services.tax_lookup import lookup_tax_code

router = APIRouter(prefix="/ledger", tags=["Module 2 · PO & Ledger"])


class SupplierIn(BaseModel):
    tax_code: str
    name: str
    supplier_code: str | None = None
    address: str = ""
    phone: str = ""
    email: str = ""


class SupplierPatch(BaseModel):
    tax_code: str | None = None
    name: str | None = None
    supplier_code: str | None = None
    address: str | None = None
    phone: str | None = None
    email: str | None = None


class POIn(BaseModel):
    po_number: str
    supplier_tax_code: str
    total_amount: float
    order_date: str = ""
    expected_delivery_date: str = ""
    employee_code: str = ""
    lines: list[dict] = []


class GRIn(BaseModel):
    receipt_number: str
    po_number: str
    received_by: str = ""
    warehouse_location: str = ""
    lines: list[dict] = []


class POUpdate(BaseModel):
    po_number: str | None = None
    supplier_tax_code: str | None = None
    total_amount: float | None = Field(default=None, ge=0)
    status: str | None = None
    order_date: str | None = None
    expected_delivery_date: str | None = None
    employee_code: str | None = None
    lines: list[dict] | None = None


class GRUpdate(BaseModel):
    receipt_number: str | None = None
    po_number: str | None = None
    received_by: str | None = None
    warehouse_location: str | None = None
    lines: list[dict] | None = None


def record_audit(db, user, action, entity_type, entity_id, detail=""):
    db.add(AuditEvent(actor_user_id=getattr(user, "id", None), actor_email=user.email, action=action, entity_type=entity_type,
                      entity_id=str(entity_id), detail=detail))


def serialize_purchase_order(row: PurchaseOrder):
    return {
        "id": row.id,
        "po_number": row.po_number,
        "supplier_tax_code": row.supplier_tax_code,
        "supplier_name": row.supplier.name if row.supplier else "",
        "total_amount": float(row.total_amount),
        "status": row.status,
        "order_date": row.order_date or "",
        "expected_delivery_date": row.expected_delivery_date or "",
        "created_by": row.created_by_user.email if row.created_by_user else "",
        "employee_code": row.employee_code or "",
        "lines": row.lines or [],
    }


def serialize_goods_receipt(row: GoodsReceipt):
    supplier = row.purchase_order.supplier if row.purchase_order else None
    return {
        "id": row.id,
        "receipt_number": row.receipt_number,
        "po_number": row.po_number,
        "supplier_name": supplier.name if supplier else "",
        "lines": row.lines or [],
        "received_at": row.received_at.isoformat(),
        "received_by": row.received_by or "",
        "warehouse_location": row.warehouse_location or "",
    }


@router.get("/tax-lookup/{tax_code}")
def tax_lookup(tax_code: str, refresh: bool = False, db: Session = Depends(get_db), _: User = Depends(current_user)):
    return lookup_tax_code(tax_code, db=db, force_live=refresh)


@router.get("/suppliers")
def suppliers(db: Session = Depends(get_db), _: User = Depends(current_user)):
    return [{"id": x.id, "supplier_code": x.supplier_code, "tax_code": x.tax_code, "name": x.name,
             "address": x.address or "", "phone": x.phone or "", "email": x.email} for x in db.query(Supplier).order_by(Supplier.name)]


@router.post("/suppliers", status_code=201)
def add_supplier(body: SupplierIn, db: Session = Depends(get_db), _: User = Depends(require_roles("admin", "chief_accountant"))):
    if db.query(Supplier).filter(Supplier.tax_code == body.tax_code).first():
        raise HTTPException(409, "Mã số thuế nhà cung cấp đã tồn tại")
    if body.supplier_code and db.query(Supplier).filter(Supplier.supplier_code == body.supplier_code).first():
        raise HTTPException(409, "Mã nhà cung cấp đã tồn tại")
    row = Supplier(**body.model_dump()); db.add(row); db.flush()
    record_audit(db, _, "supplier.create", "supplier", row.id, row.name)
    db.commit(); db.refresh(row)
    return {"id": row.id, "supplier_code": row.supplier_code, "tax_code": row.tax_code, "name": row.name,
            "address": row.address, "phone": row.phone, "email": row.email}


@router.patch("/suppliers/{supplier_id}")
def update_supplier(supplier_id: int, body: SupplierPatch, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "chief_accountant"))):
    row = db.get(Supplier, supplier_id)
    if row is None:
        raise HTTPException(404, "Không tìm thấy nhà cung cấp")
    values = body.model_dump(exclude_unset=True)
    if not values or any(value is None for value in values.values()):
        raise HTTPException(422, "Cần cung cấp giá trị cập nhật hợp lệ")
    for field in ("tax_code", "supplier_code"):
        if field in values:
            query = db.query(Supplier).filter(getattr(Supplier, field) == values[field], Supplier.id != supplier_id)
            if query.first():
                raise HTTPException(409, "Mã nhà cung cấp hoặc mã số thuế đã tồn tại")
    for key, value in values.items():
        setattr(row, key, value)
    record_audit(db, user, "supplier.update", "supplier", row.id, ",".join(sorted(values)))
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "Không thể cập nhật nhà cung cấp do trùng mã") from exc
    db.refresh(row)
    return {"id": row.id, "supplier_code": row.supplier_code, "tax_code": row.tax_code, "name": row.name,
            "address": row.address or "", "phone": row.phone or "", "email": row.email}


@router.delete("/suppliers/{supplier_id}")
def delete_supplier(supplier_id: int, db: Session = Depends(get_db), user: User = Depends(require_roles("admin", "chief_accountant"))):
    row = db.get(Supplier, supplier_id)
    if row is None:
        raise HTTPException(404, "Không tìm thấy nhà cung cấp")
    if db.query(PurchaseOrder).filter_by(supplier_id=row.id).first() or db.query(Invoice).filter_by(seller_supplier_id=row.id).first():
        raise HTTPException(409, "Không thể xóa nhà cung cấp đang được PO hoặc hóa đơn tham chiếu")
    record_audit(db, user, "supplier.delete", "supplier", row.id, row.name)
    db.delete(row)
    db.commit()
    return {"message": "Đã xóa nhà cung cấp"}


@router.get("/purchase-orders")
def purchase_orders(db: Session = Depends(get_db), _: User = Depends(current_user)):
    return [serialize_purchase_order(x) for x in db.query(PurchaseOrder).order_by(PurchaseOrder.created_at.desc())]


@router.post("/purchase-orders", status_code=201)
def add_purchase_order(body: POIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    if db.query(PurchaseOrder).filter(PurchaseOrder.po_number == body.po_number).first():
        raise HTTPException(409, "Số PO đã tồn tại")
    supplier = db.query(Supplier).filter_by(tax_code=body.supplier_tax_code).first()
    if supplier is None:
        raise HTTPException(404, "Không tìm thấy nhà cung cấp theo mã số thuế")
    row = PurchaseOrder(
        po_number=body.po_number,
        supplier_id=supplier.id,
        supplier_tax_code=supplier.tax_code,
        total_amount=body.total_amount,
        order_date=body.order_date,
        expected_delivery_date=body.expected_delivery_date,
        created_by_user_id=getattr(user, "id", None),
        employee_code=body.employee_code,
    )
    db.add(row); db.flush()
    try:
        sync_purchase_order_lines(row, body.lines)
    except ValueError as exc:
        db.rollback()
        raise HTTPException(422, str(exc)) from exc
    record_audit(db, user, "purchase_order.create", "purchase_order", row.id, row.po_number)
    db.commit(); db.refresh(row)
    return serialize_purchase_order(row)


@router.patch("/purchase-orders/{po_id}")
def update_purchase_order(po_id: int, body: POUpdate, db: Session = Depends(get_db), user: User = Depends(current_user)):
    row = db.get(PurchaseOrder, po_id)
    if not row:
        raise HTTPException(404, "Không tìm thấy PO")
    values = body.model_dump(exclude_unset=True)
    if not values or any(value is None for value in values.values()):
        raise HTTPException(422, "Cần cung cấp các giá trị cập nhật hợp lệ")
    if "po_number" in values and db.query(PurchaseOrder).filter(PurchaseOrder.po_number == values["po_number"], PurchaseOrder.id != po_id).first():
        raise HTTPException(409, "Số PO đã tồn tại")
    supplier_tax_code = values.pop("supplier_tax_code", None)
    if supplier_tax_code is not None:
        supplier = db.query(Supplier).filter_by(tax_code=supplier_tax_code).first()
        if supplier is None:
            raise HTTPException(404, "Không tìm thấy nhà cung cấp theo mã số thuế")
        row.supplier_id = supplier.id
        row.supplier_tax_code = supplier.tax_code
    lines = values.pop("lines", None)
    old_po_number = row.po_number
    for key, value in values.items():
        setattr(row, key, value)
    if lines is not None:
        try:
            sync_purchase_order_lines(row, lines)
        except ValueError as exc:
            db.rollback()
            raise HTTPException(422, str(exc)) from exc
    if row.po_number != old_po_number:
        for receipt in row.goods_receipts:
            receipt.po_number = row.po_number
        for invoice in row.invoices:
            invoice.po_number = row.po_number
            invoice.po_reference_text = row.po_number
    record_audit(db, user, "purchase_order.update", "purchase_order", row.id, ",".join(sorted(values)))
    db.commit(); db.refresh(row)
    return serialize_purchase_order(row)


@router.delete("/purchase-orders/{po_id}")
def delete_purchase_order(po_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    row = db.get(PurchaseOrder, po_id)
    if not row:
        raise HTTPException(404, "Không tìm thấy PO")
    if db.query(GoodsReceipt).filter_by(po_id=row.id).first() or db.query(Invoice).filter_by(po_id=row.id).first():
        raise HTTPException(409, "Không thể xóa PO đang được phiếu nhập hoặc hóa đơn tham chiếu")
    record_audit(db, user, "purchase_order.delete", "purchase_order", row.id, row.po_number)
    db.delete(row); db.commit()
    return {"message": "Đã xóa PO"}


@router.get("/goods-receipts")
def goods_receipts(db: Session = Depends(get_db), _: User = Depends(current_user)):
    return [serialize_goods_receipt(x) for x in db.query(GoodsReceipt).order_by(GoodsReceipt.received_at.desc())]


@router.post("/goods-receipts", status_code=201)
def add_goods_receipt(body: GRIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    purchase_order = db.query(PurchaseOrder).filter(PurchaseOrder.po_number == body.po_number).first()
    if not purchase_order:
        raise HTTPException(404, "Không tìm thấy PO")
    if db.query(GoodsReceipt).filter(GoodsReceipt.receipt_number == body.receipt_number).first():
        raise HTTPException(409, "Số phiếu nhập đã tồn tại")
    row = GoodsReceipt(
        receipt_number=body.receipt_number,
        po_id=purchase_order.id,
        po_number=purchase_order.po_number,
        purchase_order=purchase_order,
        received_by=body.received_by,
        warehouse_location=body.warehouse_location,
    )
    db.add(row); db.flush()
    try:
        sync_goods_receipt_lines(row, body.lines)
    except ValueError as exc:
        db.rollback()
        raise HTTPException(422, str(exc)) from exc
    record_audit(db, user, "goods_receipt.create", "goods_receipt", row.id, row.receipt_number)
    db.commit(); db.refresh(row)
    return serialize_goods_receipt(row)


@router.patch("/goods-receipts/{receipt_id}")
def update_goods_receipt(receipt_id: int, body: GRUpdate, db: Session = Depends(get_db), user: User = Depends(current_user)):
    row = db.get(GoodsReceipt, receipt_id)
    if not row:
        raise HTTPException(404, "Không tìm thấy phiếu nhập")
    values = body.model_dump(exclude_unset=True)
    if not values or any(value is None for value in values.values()):
        raise HTTPException(422, "Cần cung cấp các giá trị cập nhật hợp lệ")
    purchase_order = None
    if "po_number" in values:
        purchase_order = db.query(PurchaseOrder).filter_by(po_number=values.pop("po_number")).first()
        if not purchase_order:
            raise HTTPException(404, "Không tìm thấy PO")
    if "receipt_number" in values and db.query(GoodsReceipt).filter(GoodsReceipt.receipt_number == values["receipt_number"], GoodsReceipt.id != receipt_id).first():
        raise HTTPException(409, "Số phiếu nhập đã tồn tại")
    lines = values.pop("lines", None)
    for key, value in values.items():
        setattr(row, key, value)
    if purchase_order is not None:
        row.purchase_order = purchase_order
        row.po_id = purchase_order.id
        row.po_number = purchase_order.po_number
    if lines is not None:
        try:
            sync_goods_receipt_lines(row, lines)
        except ValueError as exc:
            db.rollback()
            raise HTTPException(422, str(exc)) from exc
    record_audit(db, user, "goods_receipt.update", "goods_receipt", row.id, ",".join(sorted(values)))
    db.commit(); db.refresh(row)
    return serialize_goods_receipt(row)


@router.delete("/goods-receipts/{receipt_id}")
def delete_goods_receipt(receipt_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    row = db.get(GoodsReceipt, receipt_id)
    if not row:
        raise HTTPException(404, "Không tìm thấy phiếu nhập")
    record_audit(db, user, "goods_receipt.delete", "goods_receipt", row.id, row.receipt_number)
    db.delete(row); db.commit()
    return {"message": "Đã xóa phiếu nhập"}


@router.post("/import/{kind}")
async def import_ledger(kind: str, file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(current_user)):
    if kind not in {"suppliers", "purchase-orders", "goods-receipts"}:
        raise HTTPException(404, "Loại dữ liệu không hỗ trợ")
    raw = await file.read()
    if file.filename.lower().endswith(".csv"):
        rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    elif file.filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = wb.active
        headers = [str(c or "").strip() for c in next(sheet.iter_rows(values_only=True))]
        rows = [dict(zip(headers, r)) for r in sheet.iter_rows(min_row=2, values_only=True)]
    else:
        raise HTTPException(415, "Chỉ nhận CSV hoặc XLSX")
    created, skipped = 0, 0
    for row in rows:
        try:
            if kind == "suppliers":
                key = str(row.get("tax_code", "")).strip()
                if not key or db.query(Supplier).filter_by(tax_code=key).first(): skipped += 1; continue
                db.add(Supplier(tax_code=key, name=str(row.get("name", "")).strip(), email=str(row.get("email", "") or "")))
            elif kind == "purchase-orders":
                key = str(row.get("po_number", "")).strip()
                if not key or db.query(PurchaseOrder).filter_by(po_number=key).first(): skipped += 1; continue
                supplier_tax_code = str(row.get("supplier_tax_code", "")).strip()
                supplier = db.query(Supplier).filter_by(tax_code=supplier_tax_code).first()
                if supplier is None: skipped += 1; continue
                raw_lines = row.get("lines_json", row.get("lines", ""))
                lines = json.loads(raw_lines) if isinstance(raw_lines, str) and raw_lines.strip() else []
                if not isinstance(lines, list): skipped += 1; continue
                po = PurchaseOrder(po_number=key, supplier_id=supplier.id, supplier_tax_code=supplier.tax_code,
                                   total_amount=float(row.get("total_amount", 0) or 0),
                                   order_date=str(row.get("order_date", "") or ""),
                                   expected_delivery_date=str(row.get("expected_delivery_date", "") or ""),
                                   employee_code=str(row.get("employee_code", "") or ""),
                                   created_by_user_id=getattr(user, "id", None))
                sync_purchase_order_lines(po, lines)
                db.add(po)
            else:
                key = str(row.get("receipt_number", "")).strip()
                if not key or db.query(GoodsReceipt).filter_by(receipt_number=key).first(): skipped += 1; continue
                po_number = str(row.get("po_number", "")).strip()
                purchase_order = db.query(PurchaseOrder).filter_by(po_number=po_number).first()
                if not purchase_order: skipped += 1; continue
                raw_lines = row.get("lines_json", row.get("lines", ""))
                lines = json.loads(raw_lines) if isinstance(raw_lines, str) and raw_lines.strip() else []
                if not isinstance(lines, list): skipped += 1; continue
                receipt = GoodsReceipt(receipt_number=key, po_number=po_number, po_id=purchase_order.id,
                                       purchase_order=purchase_order,
                                       received_by=str(row.get("received_by", "") or ""),
                                       warehouse_location=str(row.get("warehouse_location", "") or ""))
                sync_goods_receipt_lines(receipt, lines)
                db.add(receipt)
            created += 1
        except (TypeError, ValueError, json.JSONDecodeError):
            skipped += 1
    record_audit(db, user, "ledger.import", kind, "bulk", f"created={created};skipped={skipped}")
    db.commit()
    return {"created": created, "skipped": skipped, "total_rows": len(rows)}
